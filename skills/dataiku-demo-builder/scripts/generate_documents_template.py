#!/usr/bin/env python3
"""Synthetic documents for the RAG part of a Dataiku demo.

Produces a small corpus in THREE formats: at least one PDF, one DOCX and one
Markdown file (swap the two non-PDF formats for HTML, TXT or PPTX if it suits the
use case better). Adapt the DOCUMENTS list to the use case: the content must be
coherent with the generated datasets (same entities, rules, thresholds, vocabulary)
and in the demo language. Documents are strictly synthetic.

Local run:
    python generate_documents_template.py --out ./docs
Then upload each file to the project managed folder with the plugin tool
`upload_file_to_managed_folder`, and run the Embed recipe on that folder.

Requires: reportlab, python-docx (pip install --break-system-packages reportlab python-docx).
"""
import argparse
import os

SYNTHETIC_NOTE = "Synthetic document generated for a demonstration. Not real."

# ADAPT: one entry per document. "format" is one of: pdf, docx, md.
# Keep 8 to 20 documents in total, each 150 to 600 words, with concrete rules,
# thresholds, steps and examples that the prompt or the agents can quote.
DOCUMENTS = [
    {
        "format": "pdf",
        "name": "procedure_incident_handling",
        "title": "Incident handling procedure",
        "sections": [
            ("Purpose", "Describe how a high-severity incident is qualified and escalated."),
            ("Steps", "1. Qualify the severity. 2. Check the history of the entity. 3. Apply the escalation rule below. 4. Record the decision."),
            ("Escalation rule", "Escalate to the manager when the risk score exceeds 0.7 or when three incidents occurred in 30 days."),
        ],
    },
    {
        "format": "docx",
        "name": "policy_service_levels",
        "title": "Service level policy",
        "sections": [
            ("Scope", "This policy applies to every entity under contract."),
            ("Targets", "Critical cases are handled within 24 hours, standard cases within 5 working days."),
            ("Exceptions", "Planned maintenance windows are excluded from the targets."),
        ],
    },
    {
        "format": "md",
        "name": "playbook_recommendations",
        "title": "Recommendation playbook",
        "sections": [
            ("When risk is high", "Plan an intervention within two weeks and notify the owner."),
            ("When risk is medium", "Monitor weekly and re-score after the next event."),
            ("When risk is low", "No action; review at the next quarterly check."),
        ],
    },
]


def write_pdf(path, doc):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [Paragraph(doc["title"], styles["Title"]), Spacer(1, 12)]
    for heading, body in doc["sections"]:
        story += [Paragraph(heading, styles["Heading2"]), Paragraph(body, styles["BodyText"]), Spacer(1, 8)]
    story += [Spacer(1, 12), Paragraph(SYNTHETIC_NOTE, styles["Italic"])]
    SimpleDocTemplate(path, pagesize=A4, title=doc["title"]).build(story)


def write_docx(path, doc):
    from docx import Document

    d = Document()
    d.add_heading(doc["title"], level=0)
    for heading, body in doc["sections"]:
        d.add_heading(heading, level=1)
        d.add_paragraph(body)
    d.add_paragraph().add_run(SYNTHETIC_NOTE).italic = True
    d.save(path)


def write_md(path, doc):
    lines = [f"# {doc['title']}", ""]
    for heading, body in doc["sections"]:
        lines += [f"## {heading}", "", body, ""]
    lines += [f"_{SYNTHETIC_NOTE}_", ""]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


WRITERS = {"pdf": write_pdf, "docx": write_docx, "md": write_md}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="./docs")
    args = parser.parse_args()
    os.makedirs(args.out, exist_ok=True)
    formats = set()
    for doc in DOCUMENTS:
        path = os.path.join(args.out, f"{doc['name']}.{doc['format']}")
        WRITERS[doc["format"]](path, doc)
        formats.add(doc["format"])
        print("wrote", path)
    assert "pdf" in formats and len(formats) >= 3, "need one PDF and at least two other formats"


if __name__ == "__main__":
    main()
